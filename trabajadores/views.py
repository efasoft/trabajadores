# views.py
from django.shortcuts import render, redirect, get_object_or_404
from django.http import HttpResponse, JsonResponse
from django.conf import settings
from django.contrib.auth.decorators import login_required
from .models import Trabajador, Provincia, Ciudad
from .forms import TrabajadorForm, CiudadForm, ProvinciaForm
from datetime import datetime
from django.contrib import messages
from django.urls import reverse
import io
import os
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
from reportlab.lib import colors
from reportlab.lib.units import cm
import csv
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


from django.views.decorators.csrf import csrf_exempt
from django.utils.decorators import method_decorator
from django.views import View
import json

@csrf_exempt
def api_ciudades(request):
    if request.method == 'GET':
        provincia_id = request.GET.get('provincia')
        if provincia_id:
            ciudades = Ciudad.objects.filter(provincia_id=provincia_id).order_by('nombre')
            data = [{'id': c.id, 'nombre': c.nombre} for c in ciudades]
            return JsonResponse(data, safe=False)
        else:
            return JsonResponse([], safe=False)
    return JsonResponse({'error': 'Método no permitido'}, status=405)
    

@login_required
def home(request):
    trabajadores = Trabajador.objects.all()
    total = trabajadores.count()
    activos = trabajadores.filter(activo=True).count()
    inactivos = trabajadores.filter(activo=False).count()
    promedio = round(sum(t.sueldo_bruto for t in trabajadores) / total, 2) if total else 0

    context = {
        'total': total,
        'activos': activos,
        'inactivos': inactivos,
        'promedio': promedio
    }
    return render(request, 'trabajadores/home.html', context)


@login_required
def crear_trabajador(request):
    if request.method == 'POST':
        form = TrabajadorForm(request.POST, request.FILES)
        if form.is_valid():
            form.save()
            return JsonResponse({
                'success': True,
                'message': 'Trabajador creado correctamente.'
            })
        else:
            errors = []
            for field in form:
                for error in field.errors:
                    errors.append(f"{field.label}: {error}")
            return JsonResponse({
                'success': False,
                'errors': errors
            })

    form = TrabajadorForm()
    return render(request, 'trabajadores/form.html', {
        'form': form,
        'accion': 'Crear',
        'provincias': Provincia.objects.all(),
        'ciudades': Ciudad.objects.all()
    })


@login_required
def editar_trabajador(request, trabajador_id):
    trabajador = get_object_or_404(Trabajador, pk=trabajador_id)

    if request.method == 'POST':
        form = TrabajadorForm(request.POST, request.FILES, instance=trabajador)
        if form.is_valid():
            print("¿Tiene save?", hasattr(form, 'save'))
            print("Métodos:", [m for m in dir(form) if 'save' in m])            
            form.save()  # ✅ Ahora debería funcionar
            return JsonResponse({
                'success': True,
                'message': 'Trabajador actualizado correctamente.'
            })
        else:
            print("Errores del formulario:", form.errors)  # 🔍 Depuración            
            errors = []
            for field in form:
                for error in field.errors:
                    errors.append(f"{field.label}: {error}")
            return JsonResponse({
                'success': False,
                'errors': errors
            })

    form = TrabajadorForm(instance=trabajador)
    return render(request, 'trabajadores/form.html', {
        'form': form,
        'accion': 'Editar',
        'trabajador': trabajador,
        'provincias': Provincia.objects.all(),
        'ciudades': Ciudad.objects.all()  # ✅ Necesario para el script
    })


@login_required
def listar_trabajadores(request):
    trabajadores = Trabajador.objects.all()
    return render(request, 'trabajadores/lista.html', {'trabajadores': trabajadores})


@login_required
def eliminar_trabajador(request, trabajador_id):
    try:
        trabajador = Trabajador.objects.get(pk=trabajador_id)
        trabajador.delete()
        messages.success(request, "Trabajador eliminado correctamente.")
    except Trabajador.DoesNotExist:
        messages.error(request, "El trabajador no existe.")
    return redirect('trabajadores:listar')


@login_required
def exportar_excel(request):
    filtro = request.GET.get('q', '').lower()

    trabajadores = Trabajador.objects.all()
    if filtro:
        trabajadores = trabajadores.filter(
            nombres__icontains=filtro
        ) | trabajadores.filter(
            apellidos__icontains=filtro
        ) | trabajadores.filter(
            email__icontains=filtro
        )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Trabajadores"

    encabezados = [
        "NOMBRES", "APELLIDOS", "EMAIL", "ACTIVO", "SUELDO BRUTO",
        "FECHA", "EDAD", "TEL. CASA", "TEL. MÓVIL"
    ]

    bold_font = Font(bold=True, color="FFFFFF")
    fill = PatternFill(start_color="003C3C", end_color="003C3C", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin'),
    )

    for col_num, header in enumerate(encabezados, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = bold_font
        cell.fill = fill
        cell.alignment = center_align
        cell.border = border

    for row_num, t in enumerate(trabajadores, 2):
        row = [
            t.nombres,
            t.apellidos,
            t.email,
            "Sí" if t.activo else "No",
            f"{t.sueldo_bruto:,.2f}",
            t.fecha.strftime("%d/%m/%Y") if t.fecha else "",
            t.edad,
            t.telefono_casa,
            t.telefono_movil
        ]
        for col_num, valor in enumerate(row, 1):
            cell = ws.cell(row=row_num, column=col_num, value=valor)
            if col_num in [1, 2, 3]:
                cell.alignment = left_align
            elif col_num in [4, 6, 7, 8, 9]:
                cell.alignment = center_align
            elif col_num == 5:
                cell.alignment = right_align
            cell.border = border

    total_row = len(trabajadores) + 2
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=8)
    cell_total = ws.cell(row=total_row, column=1, value="TOTAL TRABAJADORES")
    cell_total.font = Font(bold=True)
    cell_total.alignment = right_align

    count_cell = ws.cell(row=total_row, column=9, value=trabajadores.count())
    count_cell.font = Font(bold=True)
    count_cell.alignment = center_align

    for i, col in enumerate(encabezados, 1):
        ws.column_dimensions[get_column_letter(i)].width = 18

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    disposition = 'inline' if request.GET.get('descargar') == '0' else 'attachment'
    response['Content-Disposition'] = f'{disposition}; filename="reporte_trabajadores.xlsx"'
    wb.save(response)
    return response


class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        canvas.Canvas.__init__(self, *args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            canvas.Canvas.showPage(self)
        canvas.Canvas.save(self)

    def draw_page_number(self, page_count):
        self.setFont("Helvetica", 8)
        self.drawRightString(
            A4[1] - 1.17 * cm,
            A4[0] - 2.5 * cm,
            f"Página :            "
        )
        self.drawRightString(
            A4[1] - 0.42 * cm,
            A4[0] - 2.5 * cm,
            f"  {self._pageNumber} de {page_count}"        
        )


@login_required
def exportar_pdf(request):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=1.5 * cm,
        rightMargin=1.5 * cm,
        topMargin=2.5 * cm,
        bottomMargin=0.95 * cm,
    )

    elementos = []
    filtro = request.GET.get('q', '').lower()
    trabajadores = Trabajador.objects.all()
    if filtro:
        trabajadores = trabajadores.filter(
            nombres__icontains=filtro
        ) | trabajadores.filter(
            apellidos__icontains=filtro
        ) | trabajadores.filter(
            email__icontains=filtro
        )
    trabajadores = trabajadores.order_by('-activo', 'apellidos', 'nombres')

    encabezado = [
        "APELLIDOS", "NOMBRES", "EMAIL", "EDAD", "FECHA",
        "TEL. CASA", "TEL. MOVIL", "SUELDO BRUTO €", "ACTIVO"
    ]
    data = [encabezado]

    estilo = TableStyle([])
    total_grupo = 0
    estado_actual = None
    total_general = 0

    for t in trabajadores:
        if estado_actual is not None and estado_actual != t.activo:
            idx = len(data)
            data.append(["", "", "", "", "", "", "", f"Total {'Activos' if estado_actual else 'Inactivos'}", f"{total_grupo:,.0f}"])
            estilo.add('BACKGROUND', (0, idx), (-1, idx), colors.HexColor("#666666"))
            estilo.add('TEXTCOLOR', (0, idx), (-1, idx), colors.white)
            estilo.add('FONTNAME', (0, idx), (-1, idx), 'Helvetica-Bold')
            estilo.add('ALIGN', (-1, idx), (-1, idx), 'RIGHT')
            estilo.add('TOPPADDING', (0, idx), (-1, idx), 6)
            estilo.add('BOTTOMPADDING', (0, idx), (-1, idx), 6)
            estilo.add('LINEABOVE', (7, idx), (-1, idx), 2, colors.HexColor("#707070"))
            estilo.add('LINEBELOW', (7, idx), (-1, idx), 2, colors.HexColor("#707070"))
            total_grupo = 0

        estado_actual = t.activo
        total_grupo += 1
        total_general += 1

        data.append([
            t.apellidos,            
            t.nombres,
            t.email,
            t.edad,
            t.fecha.strftime("%d/%m/%Y") if t.fecha else "",
            t.telefono_casa,
            t.telefono_movil,
            f"{t.sueldo_bruto:,.2f}",
            "SI" if t.activo else "NO"
        ])

    if estado_actual is not None:
        idx = len(data)
        data.append(["", "", "", "", "", "", "", f"Total {'Activo' if estado_actual else 'Inactivo'}", f"{total_grupo:,.0f}"])
        estilo.add('BACKGROUND', (0, idx), (-1, idx), colors.HexColor("#6599E6"))
        estilo.add('TEXTCOLOR', (0, idx), (-1, idx), colors.white)
        estilo.add('FONTNAME', (0, idx), (-1, idx), 'Helvetica-Bold')
        estilo.add('ALIGN', (-1, idx), (-1, idx), 'RIGHT')
        estilo.add('TOPPADDING', (0, idx), (-1, idx), 6)
        estilo.add('BOTTOMPADDING', (0, idx), (-1, idx), 6)
        estilo.add('LINEABOVE', (7, idx), (-1, idx), 2, colors.HexColor("#707070"))
        estilo.add('LINEBELOW', (7, idx), (-1, idx), 2, colors.HexColor("#707070"))

    idx = len(data)
    data.append(["", "", "", "", "", "", "", f"Total General", f"{total_general:,.0f}"])
    estilo.add('BACKGROUND', (0, idx), (-1, idx), colors.HexColor("#999999"))
    estilo.add('TEXTCOLOR', (0, idx), (-1, idx), colors.white)
    estilo.add('FONTNAME', (0, idx), (-1, idx), 'Helvetica-Bold')
    estilo.add('ALIGN', (-1, idx), (-1, idx), 'RIGHT')
    estilo.add('TOPPADDING', (0, idx), (-1, idx), 6)
    estilo.add('BOTTOMPADDING', (0, idx), (-1, idx), 6)
    estilo.add('LINEABOVE', (0, idx), (-1, idx), 2, colors.HexColor("#707070"))
    estilo.add('LINEBELOW', (0, idx), (-1, idx), 2, colors.HexColor("#707070"))

    col_widths = [6*cm, 6*cm, 5.5*cm, 1*cm, 2*cm, 2*cm, 2*cm, 3*cm, 1.5*cm]
    tabla = Table(data, repeatRows=1, colWidths=col_widths)

    estilo.add('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#999999"))
    estilo.add('TEXTCOLOR', (0, 0), (-1, 0), colors.white)
    estilo.add('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold')
    estilo.add('FONTSIZE', (0, 0), (-1, 0), 8.5)
    estilo.add('TOPPADDING', (0, 0), (-1, 0), 8)
    estilo.add('BOTTOMPADDING', (0, 0), (-1, 0), 8)
    estilo.add('LINEABOVE', (0, 0), (-1, 0), 2, colors.HexColor("#707070"))
    estilo.add('LINEBELOW', (0, 0), (-1, 0), 2, colors.HexColor("#707070"))
    estilo.add('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.whitesmoke, colors.HexColor("#EDEDED")])
    estilo.add('ALIGN', (0, 0), (2, -1), 'LEFT')
    estilo.add('ALIGN', (3, 0), (6, -2), 'CENTER')
    estilo.add('ALIGN', (7, 0), (7, 3), 'CENTER')
    estilo.add('ALIGN', (7, 1), (7, -2), 'RIGHT')
    estilo.add('ALIGN', (8, 0), (8, -2), 'CENTER')
    estilo.add('TEXTCOLOR', (0, 1), (-1, -1), colors.black)
    estilo.add('FONTNAME', (0, 1), (-1, -1), 'Helvetica')
    estilo.add('FONTSIZE', (0, 1), (-1, -1), 8.5)

    tabla.setStyle(estilo)
    elementos.append(tabla)

    def encabezado_pdf(canvas, doc):
        canvas.saveState()
        logo_path = os.path.join(settings.STATICFILES_DIRS[0], 'img', 'efa_soft.jpg')
        if os.path.exists(logo_path):
            canvas.drawImage(logo_path, 0.1 * cm, A4[0] - 2.5 * cm, width=4 * cm, height=2 * cm, preserveAspectRatio=True)
        canvas.setFont("Helvetica-Bold", 14)
        canvas.drawRightString(A4[1] - 0.4 * cm, A4[0] - 0.9 * cm, "RELACIÓN DE TRABAJADORES")
        canvas.setFont("Helvetica", 12)
        canvas.drawRightString(A4[1] - 0.4 * cm, A4[0] - 1.4 * cm, "Datos por Ubicación")
        canvas.setFont("Helvetica", 8)
        fecha_str = datetime.now().strftime('%d/%m/%Y')
        canvas.drawRightString(A4[1] - 0.45 * cm, A4[0] - 2.1 * cm, f"Emitido :   {fecha_str}")
        canvas.restoreState()

    doc.build(elementos, onFirstPage=encabezado_pdf, onLaterPages=encabezado_pdf, canvasmaker=NumberedCanvas)

    buffer.seek(0)
    response = HttpResponse(buffer, content_type='application/pdf')
    filename = "reporte_trabajadores.pdf"
    if request.GET.get('descargar') == '1':
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
    else:
        response['Content-Disposition'] = f'inline; filename="{filename}"'
    return response


@login_required
def listar_provincias(request):
    provincias = Provincia.objects.all().order_by("nombre")
    return render(request, "trabajadores/lista_provincias.html", {"provincias": provincias})


@login_required
def crear_provincia(request):
    if request.method == "POST":
        form = ProvinciaForm(request.POST)
        if form.is_valid():
            provincia = form.save()
            return JsonResponse({"success": True, "message": "Provincia creada con éxito"})
        else:
            return JsonResponse({"success": False, "errors": form.errors.get_json_data(escape_html=True)})
    else:
        form = ProvinciaForm()
    return render(request, "trabajadores/form_provincia.html", {"form": form, "accion": "Crear"})


@login_required
def editar_provincia(request, id):
    provincia = get_object_or_404(Provincia, id=id)
    if request.method == "POST":
        form = ProvinciaForm(request.POST, instance=provincia)
        if form.is_valid():
            form.save()
            return JsonResponse({"success": True, "message": "Provincia actualizada con éxito"})
        else:
            return JsonResponse({"success": False, "errors": form.errors.get_json_data(escape_html=True)})
    else:
        form = ProvinciaForm(instance=provincia)
    return render(request, "trabajadores/form_provincia.html", {"form": form, "accion": "Editar"})


@login_required
def eliminar_provincia(request, id):
    provincia = get_object_or_404(Provincia, id=id)
    provincia.delete()
    return redirect("trabajadores:listar_provincias")


@login_required
def listar_ciudades(request):
    ciudades = Ciudad.objects.select_related("provincia").all().order_by("nombre")
    return render(request, "trabajadores/lista_ciudades.html", {"ciudades": ciudades})


@login_required
def crear_ciudad(request):
    if request.method == "POST":
        form = CiudadForm(request.POST)
        if form.is_valid():
            ciudad = form.save()
            return JsonResponse({"success": True, "message": "Ciudad creada con éxito"})
        else:
            return JsonResponse({"success": False, "errors": form.errors.get_json_data(escape_html=True)})
    else:
        form = CiudadForm()
    provincias = Provincia.objects.all()
    return render(request, "trabajadores/form_ciudad.html", {"form": form, "accion": "Crear", "provincias": provincias})


@login_required
def editar_ciudad(request, id):
    ciudad = get_object_or_404(Ciudad, id=id)
    if request.method == "POST":
        form = CiudadForm(request.POST, instance=ciudad)
        if form.is_valid():
            form.save()
            return JsonResponse({"success": True, "message": "Ciudad actualizada con éxito"})
        else:
            return JsonResponse({"success": False, "errors": form.errors.get_json_data(escape_html=True)})
    else:
        form = CiudadForm(instance=ciudad)
    provincias = Provincia.objects.all()
    return render(request, "trabajadores/form_ciudad.html", {"form": form, "accion": "Editar", "provincias": provincias})


@login_required
def eliminar_ciudad(request, id):
    ciudad = get_object_or_404(Ciudad, id=id)
    ciudad.delete()
    return redirect("trabajadores:listar_ciudades")


